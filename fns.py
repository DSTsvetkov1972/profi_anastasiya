from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill

def format_res_file(res_file):
    # Загружаем файл
    wb = load_workbook(res_file)
    ws = wb.active


    for row in range(1, ws.max_row + 1):

        for col in range(5, 15):
            cell_to_format = ws.cell(column=col, row=row)

            if cell_to_format.value:
                cell_to_format.value = cell_to_format.value.replace('.', ',')


        first_cell_row_value = ws.cell(column=1, row=row).value 

        if 'Пробная' in str(first_cell_row_value):
           
            # Объединяем ячейки от (row, start_col) до (row, end_col)
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=ws.max_column)

            ws.merge_cells(
                start_row=row+1, start_column=1,
                end_row=row+3, end_column=1)
            
            ws.merge_cells(
                start_row=row+1, start_column=2,
                end_row=row+3, end_column=2)
            
                        
            ws.merge_cells(
                start_row=row+1, start_column=4,
                end_row=row+3, end_column=4)
            
                        
            ws.merge_cells(
                start_row=row+1, start_column=14,
                end_row=row+3, end_column=14)
            
            # Раскрашиваем ячейки
            for col in range(5, 14):
                cell_to_format = ws.cell(column=col, row=row+3)
                print(cell_to_format.value, '0,' in str(cell_to_format.value))
                if not '0.' in str(cell_to_format.value) and cell_to_format.value != '-':
                    cell_to_format.fill = PatternFill(start_color='FFE599', end_color='FFE599', fill_type='solid')

                cell_to_format = ws.cell(column=col, row=row+2)
                print(cell_to_format.value, '0,' in str(cell_to_format.value))
                if not '0.' in str(cell_to_format.value) and cell_to_format.value != '-':
                    cell_to_format.fill = PatternFill(start_color='F7CAAC', end_color='F7CAAC', fill_type='solid')    

            

   
            
    poligon_number = ''
    
    #for row in range(1, ws.max_row+1):
    #    first_cell_row_value = ws.cell(column=1, row=row).value 
    #    if poligon_number == first_cell_row_value:
    #        #ws.delete_rows(row+1)
    #        wb.save(res_file)
    #    else:
    #        poligon_number = first_cell_row_value 


    ##FFE599
    #F7CAAC
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)



    
    wb.save(res_file)