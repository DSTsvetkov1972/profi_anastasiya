from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
import pandas as pd


def permissible_concentration(fraction, pH, element, value):
    mpc = {
        'песок/суспесь': {
            'pH': None, 
            'benz': 0.02,
            'Cu': 33,
            'Zn': 55,
            'Pb': 32,
            'Cd': 0.5,
            'Ni': 20,
            'As': 2,
            'Hg': 2.1,
            'oil': 1000,     
        },
        'суглинок, pH<=5,5': {
            'pH': None, 
            'benz': 0.02,
            'Cu': 66,
            'Zn': 110,
            'Pb': 65,
            'Cd': 1,
            'Ni': 40,
            'As': 5,
            'Hg': 2.1,
            'oil': 1000,     
        },
        'суглинок, pH>5,5': {
            'pH': None, 
            'benz': 0.02,
            'Cu': 132,
            'Zn': 220,
            'Pb': 130,
            'Cd': 2.0,
            'Ni': 80,
            'As': 10,
            'Hg': 2.1,
            'oil': 1000,     
        }
    }

    if '<' in value:
        return '-'
    elif 'песок' in fraction or 'супесь' in fraction:
        return f"{float(value.replace(',', '.'))/mpc['песок/суспесь'].get(element):.2f}"
           
    elif 'суглинок' in fraction and float(pH.replace(',', '.'))<=5.5:
        return f"{float(value.replace(',', '.'))/mpc['суглинок, pH<=5,5'].get(element):.2f}"

    elif 'суглинок' in fraction and float(pH.replace(',', '.'))>5.5:
        return f"{float(value.replace(',', '.'))/mpc['суглинок, pH>5,5'].get(element):.2f}"


def get_background_res(element, value):    
    background_res_dict = {
        'pH': None, 
        'benz': None,
        'Cu': 18,
        'Zn': 43.10,
        'Pb': 19.11,
        'Cd': 0.17,
        'Ni': 15.3,
        'As': 2.62,
        'Hg': 0.03,
        'oil': None,     
        }

    value = str(value)
    if '<' in value:
        return '-'
    else:
        background_res = background_res_dict.get(element)
        res = float(value.replace(',', '.'))/background_res
        return f"{res:.2f}"


def get_source():

    df = pd.read_excel('исходник.xlsx', dtype=str, skiprows=12, header=None)

    df = df.rename(
        columns = {
            0: 'test_number',
            1: 'poligon_number',
            2: 'sample_number',
            3: 'pH',
            4: 'benz',
            5: 'Cu',
            6: 'Zn',
            7: 'Pb',
            8: 'Cd',
            9: 'Ni',
            10: 'As',
            11: 'Hg',
            12: 'oil',
            13: 'fraction'
            }
        )


    df['depth'] = df.apply(lambda row: row['test_number'] if row['poligon_number']!=row['poligon_number'] else None, axis=1)
    df['depth']= df['depth'].ffill()
    #df['fraction']= df.apply(lambda row: f"{row['depth'].replace('Глубина отбора образцов, м: ', '')}\n({row['fraction']})" , axis=1)
    df = df[df['poligon_number']== df['poligon_number']]

    return df


def get_Zc(res_row_4):
    values = []
    for k, v in res_row_4.items():
        if v == '-':
            continue
        
        if k in range(5, 12) and float(v)>1:
            values.append(float(v))
          

    return (sum(values) - (len(values)-1))        


def get_result(source_df):

    res_list = []
    for sr in source_df.itertuples():

          
        res_row_1 = {
            0: f"Пробная площадка № {sr.poligon_number} (ПП{sr.poligon_number})",
            1: None,
            2: None,
            3: None,
            4: None,
            5: None,
            6: None,
            7: None,
            8: None,
            9: None,
            10: None,
            11: None,
            12: None,
            13: None,
            'row_number': 1,            
            'sort_1': int(sr.poligon_number),
            'sort_2': int(sr.test_number),
            'sort_3': 1
        }


        res_row_3 = {
            0: None,
            1: None,
            2: 'Сi,мг/ПДК',
            3: None,
            4: permissible_concentration(sr.fraction, sr.pH, 'benz', sr.benz),
            5: permissible_concentration(sr.fraction, sr.pH, 'Cu', sr.Cu),
            6: permissible_concentration(sr.fraction, sr.pH, 'Zn', sr.Zn),
            7: permissible_concentration(sr.fraction, sr.pH, 'Pb', sr.Pb),
            8: permissible_concentration(sr.fraction, sr.pH, 'Cd', sr.Cd),
            9: permissible_concentration(sr.fraction, sr.pH, 'Ni', sr.Ni),
            10: permissible_concentration(sr.fraction, sr.pH, 'As', sr.As),
            11: permissible_concentration(sr.fraction, sr.pH, 'Hg', sr.Hg),
            12: permissible_concentration(sr.fraction, sr.pH, 'oil', sr.oil),
            13: None,
            'row_number': 3,            
            'sort_1': int(sr.poligon_number),
            'sort_2': int(sr.test_number),
            'sort_3': 3            
            }

        res_row_4 = {
            0: None,
            1: None,
            2: 'Сi,мг/фон',
            3: None,
            4: '-',
            5: get_background_res('Cu', sr.Cu),
            6: get_background_res('Zn', sr.Zn),
            7: get_background_res('Pb', sr.Pb),
            8: get_background_res('Cd', sr.Cd),
            9: get_background_res('Ni', sr.Ni),
            10: get_background_res('As', sr.As),
            11: get_background_res('Hg', sr.Hg),
            12: '-',
            13: None,
            'row_number': 4,            
            'sort_1': int(sr.poligon_number),
            'sort_2': int(sr.test_number),
            'sort_3': 4            
        }

        
        res_row_2 = {
            0: sr.test_number,
            1: f"{sr.depth.replace('Глубина отбора образцов, м: ', '')}\n({sr.fraction})",
            2: 'Сi,мг/кг',
            3: sr.pH,
            4: sr.benz,
            5: sr.Cu,
            6: sr.Zn,
            7: sr.Pb,
            8: sr.Cd,
            9: sr.Ni,
            10: sr.As,
            11: sr.Hg,
            12: sr.oil,
            13: f"{get_Zc(res_row_4):.2f}",
            'row_number': 2,
            'sort_1': int(sr.poligon_number),
            'sort_2': int(sr.test_number),
            'sort_3': 2            
        }

        df = pd.DataFrame([res_row_1, res_row_2, res_row_3, res_row_4])


        res_list.append(df)
    
    res_df = pd.concat(res_list)

    res_df = res_df.sort_values(by=['sort_1', 'sort_2', 'sort_3'])

    res_df = res_df[[0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 'row_number']]

    res_df =res_df.reset_index(drop=True)

    res_df['номер вхождения'] = res_df.groupby(0).cumcount()

    res_df = res_df[((res_df['row_number']==1)&(res_df['номер вхождения'] == 0))|((res_df['row_number']!=1))]

    return res_df   


def format_res_file(res_file):
    # Загружаем файл
    wb = load_workbook(res_file)
    ws = wb.active


    for row in range(1, ws.max_row + 1):

        for col in range(5, 15):
            cell_to_format = ws.cell(column=col, row=row)

            if cell_to_format.value:
                cell_to_format.value = cell_to_format.value.replace('.', ',')


        first_cell_row_value = ws.cell(column=15, row=row).value 

        if first_cell_row_value==1:
           
            # Объединяем ячейки от (row, start_col) до (row, end_col)
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row, end_column=ws.max_column-2)
        
        if first_cell_row_value==2:
            ws.merge_cells(
                start_row=row, start_column=1,
                end_row=row+2, end_column=1)
            
            ws.merge_cells(
                start_row=row, start_column=2,
                end_row=row+2, end_column=2)
            
                        
            ws.merge_cells(
                start_row=row, start_column=4,
                end_row=row+2, end_column=4)
            
                        
            ws.merge_cells(
                start_row=row, start_column=14,
                end_row=row+2, end_column=14)
            
            # Раскрашиваем ячейки
            for col in range(5, 14):
                cell_to_format = ws.cell(column=col, row=row+2)
                if not '0.' in str(cell_to_format.value) and cell_to_format.value != '-':
                    cell_to_format.fill = PatternFill(start_color='FFE599', end_color='FFE599', fill_type='solid')

                cell_to_format = ws.cell(column=col, row=row+1)
                if not '0.' in str(cell_to_format.value) and cell_to_format.value != '-':
                    cell_to_format.fill = PatternFill(start_color='F7CAAC', end_color='F7CAAC', fill_type='solid')    

            

   
            
    poligon_number = ''
    
    #for row in range(1, ws.max_row+1):
    #    first_cell_row_value = ws.cell(column=1, row=row).value 
    #    if poligon_number == first_cell_row_value:
    #        #ws.delete_rows(row+1)
    #        wb.save(res_file)
    #    else:
    #        poligon_number = first_cell_row_value 


    ##FFE599
    #F7CAAC
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


    ws.delete_cols(14, 2)
    
    wb.save(res_file)